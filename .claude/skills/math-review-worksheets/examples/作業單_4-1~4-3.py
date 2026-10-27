# 4-1～4-3（除法：四位÷一位、除數是二位數、末尾有 0）作業單：照「題目卷_2-1~2-2」的格式（老師 2026-10-27）
# 參考南一題庫「題目卷(1)_261027154937」（題型順序是直式→應用→填填看→選擇、計算空間太小）→ 改成
# 正面：選擇 → 填填看 → 計算 → 應用（題庫式雙欄、不給格子和提示、留足計算空間）；第二面：找錯誤、改正確（4 題）
import sys
sys.path.insert(0, r"C:\Users\class\Documents\teaching-materials\.claude\skills\math-review-worksheets\scripts")
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_TAB_ALIGNMENT
from lib import *
import lib

OUT = sys.argv[1]
ANS = []
begin("", base=13)
s = lib.doc.sections[0]
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.5)


def two_cols(sec, n):
    cols = sec._sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    c.set(qn("w:num"), str(n))
    if n == 2: c.set(qn("w:space"), "425"); c.set(qn("w:sep"), "1")
    else:
        for a in ("w:sep", "w:space"):
            if c.get(qn(a)) is not None: del c.attrib[qn(a)]
    if not cols: sec._sectPr.append(c)


def sec_title(t):
    para((t, "b"), size=15, before=6, after=3, keep=True)


def item(n, *parts, before=3, after=2, keep=False):
    p = para((f"{n}.　", "n"), *parts, before=before, after=after, keep=keep)
    p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
    return p


def space(n):
    for _ in range(n): spacer(12)


# ── 抬頭 ──
p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
run(p, "四年級數學 4-1~4-3複習卷", bold=True, size=13)
run(p, "\t____年____班　座號：_____　姓名：___________", size=12)
sec = lib.doc.add_section(WD_SECTION.CONTINUOUS); two_cols(sec, 2)

# 一、選擇題
sec_title("一、選擇題：")
CH = [("小宜有 97 張小卡片和 35 張貼紙，每 12 張小卡片包成一包，每包貼 1 張貼紙。小宜的小卡片最多可以包成幾包？",
       ["2 包", "3 包", "8 包", "23 包"]),
      ("□05÷12 的商是二位數，□ 裡不能填哪一個數字？", ["1", "2", "3", "4"]),
      ("463÷58 的商是多少？", ["6", "7", "8", "9"])]
assert 97 // 12 == 8 and 105 // 12 < 10 and 205 // 12 >= 10 and 463 // 58 == 7 and 58 * 8 > 463
for k, (stem, opts) in enumerate(CH, 1):
    p = para(("（　　）" + f"{k}. ", "n"), stem, before=3, after=0, keep=True)
    p.paragraph_format.left_indent = Cm(2.2); p.paragraph_format.first_line_indent = Cm(-2.2)
    para("　".join("①②③④"[i] + o for i, o in enumerate(opts)), indent=2.2, after=3)
ANS.append("一、1. ③（97÷12＝8…1，貼紙夠用）　2. ①（105÷12 的商是一位數）　3. ②（58×8＝464 比 463 大，商是 7）")

# 二、填填看
sec_title("二、填填看：")
item(1, "9÷5＝（　　）…（　　）\n90÷50＝（　　）…（　　）")
item(2, "6÷4＝（　　）…（　　）\n60÷40＝（　　）…（　　）")
item(3, "800÷20＝（　　）\n8000÷20＝（　　）")
item(4, "（　　　）÷12＝46…7")
assert 90 // 50 == 1 and 90 % 50 == 40 and 60 // 40 == 1 and 60 % 40 == 20 and 800 // 20 == 40 and 8000 // 20 == 400
assert 12 * 46 + 7 == 559
ANS.append("二、1. 1，4；1，40　2. 1，2；1，20　3. 40，400　4. 559（12×46＋7）")

# 三、計算題（自己寫直式，留足空間）
sec_title("三、計算題：")
CALC = [(3822, 2, 9), (7096, 9, 9), (600, 25, 7), (952, 14, 8)]   # 55÷12 太簡單、拿掉，空間給應用題   # (被除數, 除數, 留白列數)
for k, (a, b, g) in enumerate(CALC, 1):
    item(k, f"{a}÷{b}＝（　　　）" + ("…（　　）" if a % b else ""), before=4, keep=True)
    space(g)
ANS.append("三、" + "　".join(f"{k}. {a // b}" + (f"…{a % b}" if a % b else "") for k, (a, b, g) in enumerate(CALC, 1)))

# 四、應用題
sec_title("四、應用題：")
APP = ["麵包工廠做了 1305 個牛角麵包，每 8 個裝成 1 盒，共可裝成幾盒？還剩下幾個？",
       "哥哥買了一臺 7056 元的電腦螢幕，分 9 期付款，每期要付幾元？",
       "一條走道長 770 公尺，每塊地毯長 22 公尺，共需要幾塊地毯才能鋪滿？"]
for k, txt in enumerate(APP, 1):
    item(k, txt, before=4, keep=True)
    space(8 if k < 3 else 7)
assert 1305 // 8 == 163 and 1305 % 8 == 1 and 7056 // 9 == 784 and 7056 % 9 == 0 and 770 // 22 == 35 and 770 % 22 == 0
ANS.append("四、1. 1305÷8＝163…1，答：163 盒，剩 1 個　2. 7056÷9＝784，答：784 元　3. 770÷22＝35，答：35 塊")

# ── 第二面：找錯誤、改正確（單欄）──
sec = lib.doc.add_section(WD_SECTION.NEW_PAGE); two_cols(sec, 1)
sec.top_margin = sec.bottom_margin = Cm(1.2)
para(("五、找錯誤、改正確：", "b"), size=15, after=4)
hint("檢查：① 不夠商要寫 0　② 乘、減有沒有算錯　③ 餘數要比除數小",
     "驗算：商×除數＋餘數＝被除數，對得起來才是算對了。",
     title="小明寫了 4 題除法，幫他檢查：算錯的把錯的數字圈起來，在右邊寫出正確的算法")
REASONS = "錯在：□ 商中間的 0 沒寫　□ 餘數比除數大　□ 乘錯　□ 減錯"


def fix(n, divisor, dividend, q, work):
    globals()["q"](f"{n}.", f"小明算 {dividend}÷{divisor}，他算對了嗎？　□ 算對了　□ 算錯了", before=1)
    side([(["小明的算法（錯的數字圈起來）："], lambda c: ldiv(divisor, dividend, q=q, work=work, cw=0.75, h=0.56, size=14, parent=c)),
          (["正確的算法（算錯的才要寫）："], lambda c: ldiv(divisor, dividend, blank=len(work), cw=0.75, h=0.56, size=14, parent=c))])
    para(REASONS, size=12, indent=0.4, after=2)


# 1. 6024÷6：0、2 都不夠 6 除，商的百位、十位要寫 0，小明沒寫 → 14，正確 1004
assert 6024 // 6 == 1004 and 6024 % 6 == 0
fix(1, "6", "6024", "14  ", [("6   ", True), ("  24", False), ("  24", True), ("   0", False)])
# 2. 389÷42：商估成 8，餘數 53 比 42 大 → 正確 9…11
assert 42 * 8 == 336 and 389 - 336 == 53 and 389 // 42 == 9 and 389 % 42 == 11
fix(2, "42", "389", "  8", [("336", True), (" 53", False)])
# 3. 552÷23：23×4 算成 82 → 餘數 10，正確 24
assert 552 // 23 == 24 and 552 % 23 == 0 and 23 * 4 == 92
fix(3, "23", "552", " 24", [("46 ", True), (" 92", False), (" 82", True), (" 10", False)])
# 4. 816÷12：算對了（68）
assert 816 // 12 == 68 and 816 % 12 == 0 and 12 * 6 == 72 and 12 * 8 == 96
fix(4, "12", "816", " 68", [("72 ", True), (" 96", False), (" 96", True), ("  0", False)])
ANS.append("五、1. 算錯，商中間的 0 沒寫：6024÷6＝1004　2. 算錯，餘數比除數大（商要再大 1）：389÷42＝9…11　"
           "3. 算錯，乘錯（23×4＝92）：552÷23＝24　4. 算對了，68")

save(OUT)
print("saved", OUT)

# ── 解析卷（答案，另一個檔案）──
begin("", base=12)
p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
run(p, "四年級數學 4-1~4-3複習卷　解析（老師用）", bold=True, size=14)
for a in ANS:
    para(a, size=12, after=4)
ans_out = OUT.replace("題目卷", "解析卷")
save(ans_out)
print("saved", ans_out)
