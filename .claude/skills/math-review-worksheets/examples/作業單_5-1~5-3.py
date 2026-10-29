# 5-1～5-3（三角形：依角分類、依邊分類、全等）複習卷：照「題目卷_4-1~4-3」的格式，雙面（老師 2026-10-29）
# 審南一題庫「題目卷(1)_261029145230」（5 頁）：量角題 6 組重複、全等找對應出 3 次、選擇 7 分類圖不合邏輯（有三個邊→沒有三個角）
# → 正面：選擇 → 量量看（合併成 2 題）；背面：全等（有邊長的對應、A～F 對應、找全等）
import sys, math
sys.path.insert(0, r"C:\Users\class\Documents\teaching-materials\.claude\skills\math-review-worksheets\scripts")
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_TAB_ALIGNMENT
from lib import *
import lib
import figs

OUT = sys.argv[1]
ANS = []
begin("", base=13)
s = lib.doc.sections[0]
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.5)


def two_cols(sec, n):
    cols = sec._sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    c.set(qn("w:num"), str(n))
    if n == 2: c.set(qn("w:space"), "425"); c.set(qn("w:sep"), "1")
    else:
        for a in ("w:sep", "w:space"):
            if c.get(qn(a)) is not None: del c.attrib[qn(a)]
    if not cols: sec._sectPr.append(c)


def sec_title(t):
    para((t, "b"), size=15, before=6, after=3, keep=True)


def item(n, *parts, before=3, after=2, keep=False):
    p = para((f"{n}.　", "n"), *parts, before=before, after=after, keep=keep)
    p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
    return p


def wid(t):
    return sum(1 if ord(ch) > 255 else 0.5 for ch in t)


def tri_angles(A, B, C, side=3.6):
    """三個角是 A、B、C 度的三角形（照實際角度，可以量）：底邊 BC 長 side"""
    b = side * math.sin(math.radians(B)) / math.sin(math.radians(A))   # 邊 CA（對著 B）
    pB, pC = (0.0, 0.0), (side, 0.0)
    pA = (side - b * math.cos(math.radians(C)), b * math.sin(math.radians(C)))
    return [pA, pB, pC]


def rot(pts, deg, flip=False, dx=0, dy=0):
    r = math.radians(deg); out = []
    for x, y in pts:
        if flip: x = -x
        out.append((x * math.cos(r) - y * math.sin(r) + dx, x * math.sin(r) + y * math.cos(r) + dy))
    return out


def ext_angle_fig():
    """正三角形把底邊延長，∠ㄅ是外面的角（120°）"""
    cv = figs.Canvas(6.5, 3.6, 0.5, 0.6)
    P = [(0, 0), (2.8, 0), (1.4, 2.425)]
    for i in range(3): cv.line(P[i], P[(i + 1) % 3], 0.035)
    cv.line((2.8, 0), (4.6, 0), 0.035)
    for (a, b) in [(P[0], P[1]), (P[1], P[2]), (P[2], P[0])]:   # 小線條：3 邊一樣長
        m = ((a[0] + b[0]) / 2, (a[1] + b[1]) / 2); L = math.dist(a, b)
        u = ((b[0] - a[0]) / L, (b[1] - a[1]) / L); v = (-u[1], u[0])
        cv.line((m[0] - v[0] * 0.15, m[1] - v[1] * 0.15), (m[0] + v[0] * 0.15, m[1] + v[1] * 0.15), 0.025)
    cv.sector((2.8, 0), 0.45, 0, 120, (200, 200, 200))
    cv.text((3.35, 0.4), "ㄅ", 12)
    return cv.save()


# ── 抬頭 ──
p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
run(p, "四年級數學 5-1~5-3複習卷", bold=True, size=13)
run(p, "\t____年____班　座號：_____　姓名：___________", size=12)
sec = lib.doc.add_section(WD_SECTION.CONTINUOUS); two_cols(sec, 2)

# 一、選擇題
sec_title("一、選擇題：")
CH = [("3 個邊的長是下面哪一組時，是等腰三角形？", ["4、5、6 公分", "7、8、9 公分", "7、10、7 公分", "9、7、11 公分"], None),
      ("等腰三角形中，一樣大的兩個角叫作什麼？", ["頂角", "底角", "平角", "周角"], None),
      ("下面哪一個是直角三角形的說法？", ["3 個邊都一樣長", "有 1 個角比 90° 大", "有 1 個角是 90°", "3 個角都比 90° 小"], None),
      ("一個三角形的 3 個角是 30°、60°、90°，它是什麼三角形？", ["直角三角形", "等腰三角形", "正三角形", "鈍角三角形"], None),
      ("下圖是正三角形，把底邊延長，∠ㄅ是幾度？", ["60°", "90°", "120°", "150°"], ext_angle_fig()),
      ("等腰三角形一圈的長（周長）是 28 公分，其中一個腰長 10 公分，另外兩個邊各是幾公分？",
       ["8 公分、10 公分", "10 公分、10 公分", "7 公分、7 公分", "10 公分、18 公分"], None)]
assert 180 - 60 == 120 and 28 - 10 - 10 == 8
for k, (stem, opts, img) in enumerate(CH, 1):
    p = para(("（　　）" + f"{k}. ", "n"), stem, before=3, after=0, keep=True)
    p.paragraph_format.left_indent = Cm(2.2); p.paragraph_format.first_line_indent = Cm(-2.2)
    if img:
        q = lib.doc.add_paragraph(); q.paragraph_format.left_indent = Cm(2.2); q.paragraph_format.keep_with_next = True
        q.add_run().add_picture(img[0], width=Cm(img[1] * 0.75))
    if max(len(o) for o in opts) > 8:   # 長選項：一行一個，不斷字
        for i, o in enumerate(opts):
            para("①②③④"[i] + o, indent=2.2, after=3 if i == 3 else 0, keep=i < 3)
    elif sum(wid(o) + 2 for o in opts) > 14:   # 一行放不下（中文字算 1 格、數字半格）→ 兩行兩個
        for a in (0, 2):
            para("①②③④"[a] + opts[a] + "　" + "①②③④"[a + 1] + opts[a + 1], indent=2.2, after=3 if a else 0, keep=not a)
    else:
        para("　".join("①②③④"[i] + o for i, o in enumerate(opts)), indent=2.2, after=3)
ANS.append("一、1. ③　2. ②　3. ③　4. ①　5. ③（正三角形的角是 60°，180°－60°＝120°）　6. ①（28－10－10＝8，腰有兩個一樣長）")

# 二、量量看（題庫 6 組合併成 2 題；三角形照實際角度畫，量得出來）
sec_title("二、量量看：")
item(1, "用量角器量出每個三角形最大的角，再寫出是銳角、直角，還是鈍角三角形。", keep=True)
TR = [("ㄅ", (70, 60, 50)), ("ㄆ", (90, 35, 55)), ("ㄇ", (115, 40, 25)), ("ㄈ", (80, 50, 50))]
for r in (0, 2):   # 圖一列、答案一列：兩題的答案在同一高度（老師：圖歪掉看了不舒服）
    grid([[("img", figs.tri(tri_angles(*TR[r + i][1], 3.2), inner=TR[r + i][0], pt=12), 1.0)] for i in (0, 1)], cols=2)
    grid([[("p", "最大角（　　）度"), ("p", "（　　）三角形")] for i in (0, 1)], cols=2)
item(2, "量量看，再回答問題。", before=6, keep=True)
t3 = tri_angles(120, 30, 30, 4.4)
grid([[("img", figs.tri(t3, labels=[(0, "ㄊ", 0.35), (1, "ㄈ", 0.35), (2, "ㄉ", 0.35)], pt=12), 1.0)],
      [("p", "(1) ∠ㄈ＝（　　）度"), ("p", "(2) ∠ㄉ＝（　　）度"), ("p", "(3) ∠ㄊ＝（　　）度")]], cols=2)
para("(4) 這個三角形是（　　）三角形（填：正、等腰），也是（　　）三角形（填：銳角、直角、鈍角）。", indent=0.8, size=13)
ANS.append("二、1. ㄅ 70，銳角　ㄆ 90，直角　ㄇ 115，鈍角　ㄈ 80，銳角（量角容許 ±2°）　"
           "2. (1) 30　(2) 30　(3) 120　(4) 等腰，鈍角")

# ── 背面：全等 ──
sec = lib.doc.add_section(WD_SECTION.NEW_PAGE); two_cols(sec, 1)
sec_title("三、全等三角形：")
# 1. 有邊長：ㄅㄆㄇ（直角在ㄆ，ㄅㄆ 3、ㄆㄇ 4）≅ ㄉㄊㄈ（ㄅ→ㄉ、ㄆ→ㄊ、ㄇ→ㄈ）
T1 = [(0, 3), (0, 0), (4, 0)]
T2 = rot(T1, 90, dx=0, dy=0)
item(1, "下面兩個三角形全等，看圖回答問題。", keep=True)
grid([[("img", figs.tri(T1, labels=[(0, "ㄅ", 0.35), (1, "ㄆ", 0.35), (2, "ㄇ", 0.35)], right=1, sides=[((1, 2), "4 公分", 0.5)], pt=12), 0.8)],
      [("img", figs.tri(T2, labels=[(0, "ㄉ", 0.35), (1, "ㄊ", 0.35), (2, "ㄈ", 0.35)], right=1, sides=[((0, 1), "3 公分", 0.5)], pt=12), 0.8)]],
     cols=2)
para("(1) 邊ㄆㄇ的對應邊是邊（　　　），長（　　）公分；邊ㄅㄆ的對應邊是邊（　　　），長（　　）公分。", indent=0.8)
para("(2) ∠ㄉ的對應角是∠（　　），∠ㄈ的對應角是∠（　　）。", indent=0.8)
ANS.append("三、1. (1) ㄊㄈ，4；ㄉㄊ，3　(2) ㄅ，ㄇ")
# 2. A～F：ABC（直角在 B）≅ DEF（A→D、B→E、C→F），轉了方向
A1 = [(0, 3.4), (0, 0), (2.4, 0)]
A2 = rot(A1, 200, dx=0, dy=0)
item(2, "下面兩個三角形全等，看圖回答問題。", before=6, keep=True)
grid([[("img", figs.tri(A1, labels=[(0, "A", 0.35), (1, "B", 0.35), (2, "C", 0.35)], right=1, pt=12), 0.8)],
      [("img", figs.tri(A2, labels=[(0, "D", 0.35), (1, "E", 0.35), (2, "F", 0.35)], right=1, pt=12), 0.8)]], cols=2)
para("(1) 頂點 C 的對應點是頂點（　　）。　(2) 頂點 E 的對應點是頂點（　　）。", indent=0.8)
para("(3) 邊 AB 的對應邊是邊（　　）。　(4) ∠F 的對應角是∠（　　）。", indent=0.8)
ANS.append("2. (1) F　(2) B　(3) DE　(4) C")
# 3. 找全等：甲≅戊、乙≅己、丙≅丁（轉一轉、翻一翻）
item(3, "哪兩個三角形全等？（可以轉一轉、翻一翻）", before=6, keep=True)
J = [(0, 0), (2.6, 0), (0.6, 1.6)]          # 甲：一般三角形
Y = [(0, 0), (2.0, 0), (1.0, 1.7)]          # 乙：等腰（較胖）
B_ = [(0, 0), (1.2, 0), (0.6, 2.4)]         # 丙：等腰（瘦高）
SH = [("甲", J), ("乙", Y), ("丙", B_), ("丁", rot(B_, 140)), ("戊", rot(J, 75, flip=True)), ("己", rot(Y, 200))]
grid([[("img", figs.tri(p, inner=n, pt=12), 0.75)] for n, p in SH], cols=6)
para("(1) 甲和（　　）全等。　(2) 乙和（　　）全等。　(3) 丙和（　　）全等。", indent=0.8)
ANS.append("3. (1) 戊　(2) 己　(3) 丁")

# 四、是非（概念）
sec_title("四、對的畫○，錯的畫×：")
for k, st in enumerate(["正三角形也是等腰三角形。", "直角三角形可能有 2 個直角。", "全等的兩個三角形，對應角一樣大。",
                        "鈍角三角形可能有 2 個鈍角。"], 1):
    p = para(("（　　）" + f"{k}. ", "n"), st, before=2, after=2)
ANS.append("四、1. ○（3 個邊一樣長，當然有 2 個邊一樣長）　2. ×　3. ○　4. ×（最多 1 個鈍角）")

save(OUT)
print("saved", OUT)

begin("", base=12)
p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
run(p, "四年級數學 5-1~5-3複習卷　解析（老師用）", bold=True, size=14)
for a in ANS:
    para(a, size=12, after=4)
ans_out = OUT.replace("題目卷", "解析卷")
save(ans_out)
print("saved", ans_out)
