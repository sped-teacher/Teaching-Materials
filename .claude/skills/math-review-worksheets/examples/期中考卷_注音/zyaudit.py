# 檢查期中考裡的破音字：列出每個破音字用了哪個注音字型、在哪個詞裡，並畫出目前顯示的注音
import os, sys, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import zy
import zyfont
from docx import Document
from PIL import Image, ImageDraw, ImageFont

D = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\class\Desktop\學數4上\期中考"
FILES = ["四上數學期中考_情境版.docx", "四上數學期中考_精簡版_報讀.docx"]
seen = {}
_poly = {}


def poly(ch):
    """這個字在破音一～五字型裡，有沒有和預設字型不一樣的注音"""
    if ch not in _poly:
        base = zyfont.crop(ch, 0)
        _poly[ch] = base is not None and any(
            zyfont.crop(ch, i) is not None and zyfont.diff(base, zyfont.crop(ch, i)) > 40 for i in range(1, 6))
    return _poly[ch]
for fn in FILES:
    doc = Document(os.path.join(D, fn))
    paras = list(doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                paras += c.paragraphs
    for p in paras:
        text = "".join(r.text for r in p.runs)
        pos = 0
        for r in p.runs:
            name = r.font.name or ""
            for ch in r.text:
                if name.startswith("文鼎標楷注音") and poly(ch):
                    idx = zyfont.NAMES.index(name)
                    ctx = text[max(0, pos - 3): pos + 4]
                    seen.setdefault((ch, idx), set()).add(ctx)
                pos += len(ch)
items = sorted(seen.items(), key=lambda kv: kv[0])
print(len(items), "個（字, 字型）組合")
# 畫出來
F = 70
cols = 6
W, H = 330, 120
img = Image.new("RGB", (cols * W, ((len(items) + cols - 1) // cols) * H), "white")
dr = ImageDraw.Draw(img)
lab = ImageFont.truetype(r"C:\Windows\Fonts\kaiu.ttf", 20)
for i, ((ch, idx), ctxs) in enumerate(items):
    x, y = (i % cols) * W, (i // cols) * H
    dr.text((x + 8, y + 8), ch, font=ImageFont.truetype(zyfont.TTC, F, index=idx), fill="black")
    dr.text((x + 110, y + 10), f"字型{idx}", font=lab, fill="gray")
    ctx = sorted(ctxs)[0].replace("\n", "")
    dr.text((x + 110, y + 40), ctx, font=lab, fill="black")
    
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "zyaudit.png")
img.save(out)
for (ch, idx), ctxs in items:
    print(ch, idx, " | ".join(sorted(ctxs)[:3]))
print(out)
