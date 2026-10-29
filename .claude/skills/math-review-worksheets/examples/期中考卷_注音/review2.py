# 期中複習考 2（第 4～5 單元：除法、三角形）——照複習考 1／期中考卷格式（老師 2026-10-29）
# A4、雙欄、正反兩面、注音；不給格子和提示；有鑑別題；保留 1 題畫圖（老師 2026-10-29）。
import sys, os, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ["EXAM_ZS"] = "12"   # A4 欄比較窄：注音字 12 號
import exam as E
from exam import Z, zp, head, choice, cells, center, ck, lib, figs
from lib import *
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK

COLW = 8.6
ANS = []


def build(path):
    begin("", base=E.ZS, margin=1.2)
    s = lib.doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)      # A4
    s.left_margin = s.right_margin = Cm(1.2)
    s.top_margin = s.bottom_margin = Cm(1.0)
    p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "四年級數學　期中複習考 2　第四單元～第五單元", bold=True, size=15)
    p = lib.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, "座號：＿＿＿　姓名：＿＿＿＿＿＿＿", size=13)
    sec = lib.doc.add_section(WD_SECTION.CONTINUOUS)
    cols = sec._sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    c.set(qn("w:num"), "2"); c.set(qn("w:space"), "480"); c.set(qn("w:sep"), "1")
    if not cols: sec._sectPr.append(c)
    W = [COLW / 2] * 2
    BOT = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    # 一、是非
    hp = zp(("一、", "b"), ("對的畫○，錯的畫", "b"), size=15, before=2, after=0, keep=True)
    sr = hp.add_run("✕"); sr.font.size = Pt(17); sr.bold = True; sr.font.name = "Segoe UI Symbol"
    rf = sr._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), "Segoe UI Symbol")
    E.thicken(sr)
    zp(("每題 2 分，共 8 分", "s"), size=11, after=2, keep=True).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for k, st in enumerate(["除法的餘數一定比除數小。", "等腰三角形的 3 個邊都一樣長。", "正三角形也是銳角三角形。",
                            "形狀一樣的兩個三角形，一定全等。"], 1):
        zp(("（　）(%d) " % k, "n"), st, indent=1.6, first=-1.6, after=3)
    ANS.append("一、(1) ○　(2) ×（只有 2 個邊一樣長）　(3) ○（3 個角都是 60°）　(4) ×（大小也要一樣）")

    # 二、選擇
    head("二", "選擇題", "每題 2 分，共 10 分")
    choice(1, "382÷47 的商是多少？", ["6", "7", "8", "9"])
    choice(2, "下面哪一個算式的商是二位數？", ["245÷36", "318÷42", "512÷48", "396÷45"])
    choice(3, "5003÷5＝？", ["100…3", "1000…3", "1001", "10003"])
    def short_choice(n, stem, opts):   # A4 欄窄：短選項間距小一點、縮排往左，排在同一行
        zp(("（　）(%d) " % n, "n"), stem, indent=1.6, first=-1.6, keep=True)
        zp(*sum([[("①②③④"[i], "n"), o, (" ", "n")] for i, o in enumerate(opts)], []), indent=0.9, after=4)
    short_choice(4, "一個三角形的 3 個角是 30°、60°、90°，它是（　）三角形。", ["銳角", "直角", "鈍角", "正"])
    short_choice(5, "3 個邊是 5 公分、5 公分、3 公分的三角形，是（　）三角形。", ["正", "等腰", "直角", "鈍角"])
    ck(382 // 47 == 8 and [245 // 36, 318 // 42, 512 // 48, 396 // 45] == [6, 7, 10, 8] and 5003 // 5 == 1000 and 5003 % 5 == 3, "c")
    ANS.append("二、(1) ③（47×8＝376；47×9＝423 太大）　(2) ③（51 比 48 大，商從十位開始）　(3) ②（鑑別題）　(4) ②　(5) ②")

    # 三、看圖分類
    head("三", "看圖填代號", "每格 2 分，共 8 分")
    zp("小線條一樣多的邊一樣長，□ 是直角。", size=E.ZS, keep=True)
    shapes = [("甲", [(0, 0), (3, 0), (1.5, 2.598)], dict(marks=[((0, 1), 1), ((1, 2), 1), ((2, 0), 1)])),
              ("乙", [(0, 0), (3.2, 0), (0, 2.2)], dict(right=0)),
              ("丙", [(0, 0), (4.2, 0), (2.9, 1.0)], dict()),
              ("丁", [(0, 0), (2.6, 0), (1.3, 2.8)], dict(marks=[((0, 2), 1), ((1, 2), 1)]))]
    cells([[("img", figs.tri(p, inner=n, pt=12, **kw), 0.42)] for n, p, kw in shapes[:2]], W, valign=BOT)
    cells([[("img", figs.tri(p, inner=n, pt=12, **kw), 0.42)] for n, p, kw in shapes[2:]], W, valign=BOT)
    cells([[("z", ("(1) ", "n"), "正三角形", ("（　）", "n")), ("z", ("(3) ", "n"), "直角三角形", ("（　）", "n"))],
           [("z", ("(2) ", "n"), "等腰三角形", ("（　）", "n")), ("z", ("(4) ", "n"), "鈍角三角形", ("（　）", "n"))]], W)
    ANS.append("三、(1) 甲　(2) 丁（甲也有 2 邊一樣長，寫甲丁也給分）　(3) 乙　(4) 丙")

    # 四、填充
    head("四", "填充題", "每格 2 分，共 22 分")
    zp(("(1) 70÷30＝（　　）…（　　）", "n"), indent=0.8, first=-0.8)          # 鑑別題：餘數是 10 不是 1
    zp(("(2) 900÷30＝（　　），9000÷30＝（　　）", "n"), indent=0.8, first=-0.8)
    zp(("(3) （　　　）÷24＝15…9", "n"), indent=0.8, first=-0.8)
    zp(("(4) ", "n"), "正三角形的每一個角都是", ("（　　）", "n"), "度。", indent=0.8, first=-0.8)
    zp(("(5) ", "n"), "150 顆糖，每 24 顆裝一盒，可以裝滿", ("（　　）", "n"), "盒，還剩", ("（　　）", "n"), "顆。", indent=0.8, first=-0.8)
    zp(("(6) ", "n"), "□32÷36 的商是一位數，□ 裡最大可以填", ("（　　）", "n"), "。", indent=0.8, first=-0.8)   # 鑑別題
    zp(("(7) ", "n"), "一個三角形有 2 個角是 45°，1 個角是 90°，它是", ("（　　）", "n"), "三角形，也是",
       ("（　　）", "n"), "三角形。", indent=0.8, first=-0.8)
    ck(70 // 30 == 2 and 70 % 30 == 10 and 900 // 30 == 30 and 9000 // 30 == 300 and 24 * 15 + 9 == 369
       and 150 // 24 == 6 and 150 % 24 == 6 and 332 // 36 < 10 and 432 // 36 >= 10, "f")
    ANS.append("四、(1) 2，10（鑑別題）　(2) 30，300　(3) 369（24×15＋9）　(4) 60　(5) 6，6　(6) 3（332÷36＝9…8；432÷36 商是二位數，鑑別題）　"
               "(7) 直角，等腰（順序可以對調）")

    zp(("背面還有題目 →", "b"), size=12, before=0).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    lib.doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    # 五、直式
    head("五", "用直式算算看", "每題 5 分，共 20 分", below=True)
    blank = lambda n: (lambda c: [c.add_paragraph() for _ in range(n)])

    def ans_line(txt, n, size=13, ind=0.5):
        def f(c):
            q = c.add_paragraph(); q.paragraph_format.left_indent = Cm(ind); run(q, txt, size=size)
            blank(n)(c)
        return f
    side([(["(1) 7254÷6"], ans_line("＝（　　　　）", 7)), (["(2) 3605÷7"], ans_line("＝（　　　　）", 7))], widths=W, keep=False)
    side([(["(3) 786÷23"], ans_line("＝（　）…（　）", 6, size=12, ind=0)), (["(4) 924÷28"], ans_line("＝（　　　　）", 6))], widths=W, keep=False)
    ck(7254 // 6 == 1209 and 7254 % 6 == 0 and 3605 // 7 == 515 and 3605 % 7 == 0 and 786 // 23 == 34 and 786 % 23 == 4
       and 924 // 28 == 33 and 924 % 28 == 0, "v")
    ANS.append("五、(1) 1209（商中間有 0，鑑別題）　(2) 515　(3) 34…4　(4) 33")

    # 六、全等（課本式：左邊標角和邊，右邊轉了方向、不標）
    head("六", "兩個三角形全等", "每格 2 分，共 12 分", below=True)
    Bp, Cp = (0.0, 0.0), (5.0, 0.0)
    ab = 5.0 * math.sin(math.radians(30)) / math.sin(math.radians(100))
    Ap = (ab * math.cos(math.radians(50)), ab * math.sin(math.radians(50)))   # ∠A100° ∠B50° ∠C30°

    def Tf(p, deg=160):
        x, y = -p[0], p[1]; r = math.radians(deg)
        return (x * math.cos(r) - y * math.sin(r), x * math.sin(r) + y * math.cos(r))
    t1 = figs.tri([Ap, Bp, Cp], labels=[(0, "A", 0.35), (1, "B", 0.35), (2, "C", 0.35)],
                  angles=[(0, "100°", 0.5), (1, "50°", 0.85), (2, "30°", 1.4)],
                  sides=[((0, 1), "6 公分", 0.5), ((0, 2), "9 公分", 0.5)])
    t2 = figs.tri([Tf(Ap), Tf(Cp), Tf(Bp)], labels=[(0, "D", 0.35), (1, "E", 0.35), (2, "F", 0.35)])
    cells([[("img", t1, 0.62)], [("img", t2, 0.62)]], W, valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    zp(("(1) ", "n"), "頂點 A 的對應點是頂點", ("（　　）", "n"), "。", indent=0.8, first=-0.8)
    zp(("(2) ", "n"), "邊 DE 的對應邊是邊", ("（　　）", "n"), "，\nDE 長", ("（　　）", "n"), "公分。", indent=0.8, first=-0.8)
    zp(("(3) ", "n"), "∠E 的對應角是∠", ("（　　）", "n"), "，\n∠E 是", ("（　　）", "n"), "度。", indent=0.8, first=-0.8)
    zp(("(4) ", "n"), "三角形 DEF 是", ("（　　）", "n"), "三角形。（銳角、直角、鈍角）", indent=0.8, first=-0.8)
    ANS.append("六、(1) D　(2) AC，9　(3) C，30　(4) 鈍角（∠D＝∠A＝100°）")

    lib.doc.paragraphs[-1].add_run().add_break(WD_BREAK.COLUMN)
    # 七、畫畫看：保留 1 題（背面右欄最上面）（畫正三角形：量角器畫 60°，再量 3 公分）
    head("七", "畫畫看", "共 4 分")
    zp("以這條 3 公分的線為一邊，從頂點畫一個正三角形。", indent=0.3, keep=True)
    for _ in range(5):
        e = lib.doc.add_paragraph(); e.paragraph_format.keep_with_next = True
    image(figs.baseline(3.0))
    ANS.append("七、畫出邊長 3 公分的正三角形（在頂點畫 60°、新邊量 3 公分再連起來；邊長容許 ±0.2 公分）")

    # 八、應用題
    head("八", "應用題（記得寫「答」）", "(1)(2) 每題 5 分，(3) 6 分，共 16 分", below=True)
    probs = ["1350 公克的糖，每 8 公克包成一包，可以包幾包？還剩幾公克？",
             "學校買了 28 箱飲料，一共 672 瓶，平均一箱有幾瓶？",
             "有 925 顆糖，先拿走 25 顆，剩下的每 36 顆裝成一盒，可以裝成幾盒？"]   # (3) 兩步驟（鑑別題）
    for k, txt in enumerate(probs, start=1):
        zp(("(%d) " % k, "n"), txt, indent=0.8, first=-0.8, keep=True)
        zp(("　", "n"), indent=0.8, after=56 if k < 3 else 58)
    ck(1350 // 8 == 168 and 1350 % 8 == 6 and 672 // 28 == 24 and 672 % 28 == 0 and (925 - 25) // 36 == 25 and (925 - 25) % 36 == 0, "a")
    ANS.append("八、(1) 1350÷8＝168…6，168 包，剩 6 公克　(2) 672÷28＝24，24 瓶　"
               "(3) 925－25＝900，900÷36＝25，25 盒（鑑別題）")
    p = zp("□我有仔細讀題目　□我有盡力完成", size=12, before=4, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    save(path)


def answers(path):
    begin("", base=12)
    p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
    run(p, "四年級數學　期中複習考 2（第四～五單元）答案與評分", bold=True, size=14)
    for a in ANS:
        para(a, size=12, after=4)
    para("評分：應用題 (1)(2) 列式 2 分、答 3 分（(1) 包數、剩幾公克少寫一個扣 1 分）；(3) 兩個算式各 2 分、答 2 分"
         "（只算到 925－25＝900 給 2 分）。計算題只看答案。", size=12, before=6)
    save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    build(os.path.join(out, "複習考2_第4-5單元.docx"))
    answers(os.path.join(out, "複習考2_答案與評分.docx"))
    import zy
    if zy.warn: print("注音沒把握：", sorted(set(zy.warn)))
    print("ok")
