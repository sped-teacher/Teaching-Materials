# 2-3（三位數×二位數）作業單，程度較好的學生用：照「題目卷_2-1~2-2_加第二面」的格式
# 正面：題庫式雙欄（選擇、填填看、計算、應用題，不給格子和提示）；第二面：找錯誤、改正確（4 題）
import sys
sys.path.insert(0, r"C:\Users\class\Documents\teaching-materials\.claude\skills\math-review-worksheets\scripts")
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_TAB_ALIGNMENT
from lib import *
import lib

OUT = sys.argv[1]
ANS = []
begin("", base=13)
s = lib.doc.sections[0]
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.5)


def two_cols(sec, n):
    cols = sec._sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    c.set(qn("w:num"), str(n))
    if n == 2: c.set(qn("w:space"), "425"); c.set(qn("w:sep"), "1")
    else:
        for a in ("w:sep", "w:space"):
            if c.get(qn(a)) is not None: del c.attrib[qn(a)]
    if not cols: sec._sectPr.append(c)


def sec_title(t):
    para((t, "b"), size=15, before=6, after=3, keep=True)


def item(n, *parts, before=3, after=2, keep=False):
    p = para((f"{n}.　", "n"), *parts, before=before, after=after, keep=keep)
    p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
    return p


def stacked(n, a, b, gap=8):
    """題庫式直式：題號＋被乘數、乘數和底線，下面留空白算"""
    t = lib.doc.add_table(rows=2, cols=2); borders(t, sz=0)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell_text(t.cell(0, 0), f"{n}.", size=13)
    t.cell(0, 0).width = t.cell(1, 0).width = Cm(0.9)
    for k, txt in enumerate([str(a), "×　" + str(b)]):
        c = t.cell(k, 1); c.width = Cm(2.8)
        cell_text(c, txt, size=14)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bd = OxmlElement("w:tcBorders"); bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:color", "000000")): bt.set(qn(k), v)
    bd.append(bt); t.cell(1, 1)._tc.get_or_add_tcPr().append(bd)
    for _ in range(gap): spacer(12)


# ── 抬頭 ──
p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)   # 班級姓名靠右，不擠成兩行
run(p, "四年級數學 2-3複習卷", bold=True, size=13)
run(p, "	____年____班　座號：_____　姓名：___________", size=12)
sec = lib.doc.add_section(WD_SECTION.CONTINUOUS); two_cols(sec, 2)

# 一、選擇題
sec_title("一、選擇題：")
CH = [("已知 26×4＝104，那麼 260×40＝？", ["1040", "10400", "104000", "1004"], 2),
      ("398×21 的積，最接近下面哪一個數？", ["800", "6000", "8000", "80000"], 3),
      ("下面哪一個算式的積最大？", ["305×20", "350×20", "305×21", "300×25"], 4)]
assert 260 * 40 == 10400 and 400 * 20 == 8000
assert max(305 * 20, 350 * 20, 305 * 21, 300 * 25) == 300 * 25
for k, (stem, opts, a) in enumerate(CH, 1):
    p = para(("（　　）" + f"{k}. ", "n"), stem, before=3, after=0, keep=True)
    p.paragraph_format.left_indent = Cm(2.2); p.paragraph_format.first_line_indent = Cm(-2.2)
    para("　".join("①②③④"[i] + o for i, o in enumerate(opts)), indent=2.2, after=3)
ANS.append("一、1. ②　2. ③（400×20＝8000）　3. ④（6100、7000、6405、7500）")

# 二、填填看
sec_title("二、填填看：")
item(1, "4×6＝（　　　）\n40×6＝（　　　）\n400×60＝（　　　）")
item(2, "125×8＝（　　　）\n125×80＝（　　　）")
item(3, "28×35＝980，所以\n280×35＝（　　　）\n28×350＝（　　　）")
assert 4 * 6 == 24 and 40 * 6 == 240 and 400 * 60 == 24000 and 125 * 8 == 1000 and 125 * 80 == 10000
assert 28 * 35 == 980 and 280 * 35 == 9800 and 28 * 350 == 9800
ANS.append("二、1. 24，240，24000　2. 1000，10000　3. 9800，9800")

# 三、計算題（用直式）
sec_title("三、計算題：")
CALC = [(305, 24), (2047, 16), (860, 45), (238, 50)]   # 中間有 0、四位數×二位數、被乘數末尾 0、乘整十
for k, (a, b) in enumerate(CALC, 1):
    stacked(k, a, b)
ANS.append("三、" + "　".join(f"{k}. {a * b}" for k, (a, b) in enumerate(CALC, 1)))

# 四、應用題
sec_title("四、應用題：")
APP = ["1 箱蘋果有 128 顆，36 箱共有幾顆？",
       "學校買了 24 台腳踏車，\n每台 3250 元，共要付幾元？",   # 手動換行：數字和單位不要拆開
       "1 本故事書有 205 頁，小明讀完了 18 本，共讀了幾頁？",
       "1 盒鉛筆有 12 枝，每枝 15 元。買 24 盒鉛筆，共要付幾元？"]
for k, txt in enumerate(APP, 1):
    item(k, txt, before=4, keep=True)
    for _ in range(5 if k < 4 else 4): spacer(12)
assert 128 * 36 == 4608 and 3250 * 24 == 78000 and 205 * 18 == 3690 and 12 * 15 == 180 and 180 * 24 == 4320
ANS.append("四、1. 128×36＝4608，答：4608 顆　2. 3250×24＝78000，答：78000 元　3. 205×18＝3690，答：3690 頁　"
           "4. 12×15＝180，180×24＝4320（或 15×12×24），答：4320 元")

# ── 第二面：找錯誤、改正確（單欄）──
sec = lib.doc.add_section(WD_SECTION.NEW_PAGE); two_cols(sec, 1)
sec.top_margin = sec.bottom_margin = Cm(1.2)   # 第二面 4 題直式格要排進一頁
para(("五、找錯誤、改正確：", "b"), size=15, after=4)
hint("檢查：① 每一位乘對了嗎　② 進位加了嗎　③ 乘十位的第一個數字寫在十位（對齊）",
     "先估一估：305×24 大約是 300×24＝7200，差很多就一定算錯了。",
     title="小明寫了 4 題直式，幫他檢查：算錯的把錯的數字圈起來，在右邊寫出正確的算法")
REASONS = "錯在：□ 九九算錯　□ 忘記加進位　□ 第 2 列沒有對齊　□ 加錯"


def fix(n, expr, student_rows, blank_rows, ncols=5):
    q(f"{n}.", f"小明算 {expr}，他算對了嗎？　□ 算對了　□ 算錯了", before=1)
    side([(["小明的算法（錯的數字圈起來）："], lambda c: vgrid(ncols, student_rows, h=0.56, size=13, cw=0.75, parent=c)),
          (["正確的算法（算錯的才要寫）："], lambda c: vgrid(ncols, blank_rows, h=0.56, size=13, cw=0.75, parent=c))])
    para(REASONS, size=12, indent=0.4, after=2)


def blank(a, b):
    return [dict(d=str(a)), dict(op="×", d=str(b), line=True), dict(), dict(op="＋", line=True, start=1), dict()]


# 1. 305×24：乘 4 時，0 那一位忘記加進位（0×4＋2＝2）→ 第 1 列寫成 1200，正確 1220；答案 7300，正確 7320
assert 305 * 4 == 1220 and 305 * 24 == 7320 and 1200 + 6100 == 7300
fix(1, "305×24", [dict(d="305"), dict(op="×", d="24", line=True), dict(d="1200"), dict(op="＋", d="610 ", line=True), dict(d="7300")],
    blank(305, 24))
# 2. 476×38：乘十位的 1428 沒有從十位寫 → 3808＋1428＝5236，正確 18088
assert 476 * 8 == 3808 and 476 * 3 == 1428 and 3808 + 1428 == 5236 and 476 * 38 == 18088
fix(2, "476×38", [dict(d="476"), dict(op="×", d="38", line=True), dict(d="3808"), dict(op="＋", d="1428", line=True), dict(d="5236")],
    blank(476, 38))
# 3. 860×45：乘 4 時 6×4 算成 28 → 第 2 列寫成 3480，正確 3440；答案 39100，正確 38700
assert 860 * 4 == 3440 and 4300 + 34800 == 39100 and 860 * 45 == 38700
fix(3, "860×45", [dict(d="860"), dict(op="×", d="45", line=True), dict(d="4300"), dict(op="＋", d="3480 ", line=True), dict(d="39100")],
    blank(860, 45))
# 4. 609×17：算對了（4263＋6090＝10353）
assert 609 * 7 == 4263 and 4263 + 6090 == 10353 and 609 * 17 == 10353
fix(4, "609×17", [dict(d="609"), dict(op="×", d="17", line=True), dict(d="4263"), dict(op="＋", d="609 ", line=True), dict(d="10353")],
    blank(609, 17))
ANS.append("五、1. 算錯，忘記加進位：305×4＝1220，1220＋6100＝7320　2. 算錯，第 2 列沒有對齊：3808＋14280＝18088　"
           "3. 算錯，九九算錯（6×4＝24）：860×4＝3440，4300＋34400＝38700　4. 算對了，10353")

save(OUT)
print("saved", OUT)

# ── 解析卷（答案，另一個檔案）──
begin("", base=12)
p = lib.doc.paragraphs[0] if lib.doc.paragraphs else lib.doc.add_paragraph()
run(p, "四年級數學 2-3複習卷　解析（老師用）", bold=True, size=14)
for a in ANS:
    para(a, size=12, after=4)
ans_out = OUT.replace("題目卷", "解析卷")
save(ans_out)
print("saved", ans_out)
