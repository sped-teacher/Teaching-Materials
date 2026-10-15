# 資源班數學學習單的版面工具（A4 直式、標楷體、位值表、直式格、除法直式、圖）
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "標楷體"
WAN, GE, HOT, GRAY = "DCE9F5", "FFF4CC", "FBD5B5", "EDEDED"
EX = RGBColor(0x1F, 0x5F, 0xA8)
B2 = "（　　）"
B4 = "（　　　　）"
B5 = "（　　　　　）"
B6 = "（　　　　　　）"
ALL = ["千萬", "百萬", "十萬", "萬", "千", "百", "十", "一"]
doc = None
HEADER = ""


def begin(header, base=14, margin=1.5):
    global doc, HEADER
    HEADER = header
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(margin)
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(base)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.2
    return doc


PRE_SAVE = None


def save(path):
    if PRE_SAVE: PRE_SAVE()
    doc.save(path)


def shade(el_pr, fill):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill)
    el_pr.append(s)


def run(p, text, bold=False, size=None, color=None, block=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if block:
        rpr = r._element.get_or_add_rPr()
        b = OxmlElement("w:bdr")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "0"), ("w:color", "7F7F7F")):
            b.set(qn(k), v)
        rpr.append(b); shade(rpr, GRAY)
    return r


def _parts(p, parts, size=None):
    for x in parts:
        if isinstance(x, tuple):
            t, k = x
            run(p, t, bold=(k == "b"), size=size, color=EX if k == "ex" else None, block=(k == "blk"))
        else:
            run(p, x, size=size)


def para(*parts, size=None, indent=0, before=0, after=2, align=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(indent); pf.space_before = Pt(before); pf.space_after = Pt(after)
    if keep: pf.keep_with_next = True
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _parts(p, parts, size)
    return p


def spacer(pt=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pt)
    run(p, "", size=pt)
    return p


def page_break():
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def cell_text(c, text, bold=False, size=None, color=None, fill=None, align="c"):
    c.text = ""
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "c" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)   # 套在別人的文件裡時，不要吃到原本樣式的段前距
    p.paragraph_format.line_spacing = 1.0
    if text: run(p, text, bold=bold, size=size, color=color)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill: shade(c._tc.get_or_add_tcPr(), fill)


def borders(tbl, sz=8, color="595959"):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}")
        x.set(qn("w:val"), "single" if sz else "nil"); x.set(qn("w:sz"), str(sz)); x.set(qn("w:color"), color)
        b.append(x)
    tblPr.append(b)


def cell_border(c, **kw):
    """cell_border(c, bottom=(sz, color), top=..., right=..., left=...)；sz=0 表示不要線"""
    tcPr = c._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        if side in kw:
            sz, color = kw[side]
            x = OxmlElement(f"w:{side}")
            x.set(qn("w:val"), "single" if sz else "nil"); x.set(qn("w:sz"), str(sz)); x.set(qn("w:color"), color)
            b.append(x)


def row_h(row, cm):
    row.height = Cm(cm); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def no_split(tbl):
    for r in tbl.rows:
        trPr = r._tr.get_or_add_trPr()
        x = OxmlElement("w:cantSplit"); trPr.append(x)
        for c in r.cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True


def place_table(cols, rows, label_w=None, hot=None, ex_rows=(), cell_w=1.75, h=1.0, size=18):
    """位值表。rows：[(左欄標籤或 None, 數字字串)]；hot：{列: 位名}"""
    hot = hot or {}
    has_label = any(lab is not None for lab, _ in rows)
    off = 1 if has_label else 0
    t = doc.add_table(rows=1 + len(rows), cols=len(cols) + off)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    if has_label: cell_text(t.cell(0, 0), "")
    for j, c in enumerate(cols):
        cell_text(t.cell(0, off + j), c + "位", bold=True, size=11, fill=WAN if "萬" in c else GE)
    row_h(t.rows[0], 0.75)
    for i, (lab, digits) in enumerate(rows, start=1):
        if has_label: cell_text(t.cell(i, 0), lab or "", size=12, bold=True)
        d = digits.rjust(len(cols)) if digits else ""
        for j, c in enumerate(cols):
            ch = d[j].strip() if d else ""
            cell_text(t.cell(i, off + j), ch, size=size, color=EX if (i - 1) in ex_rows else None,
                      fill=HOT if hot.get(i - 1) == c else None)
        row_h(t.rows[i], h)
    for r in t.rows:
        if has_label: r.cells[0].width = Cm(label_w or 2.4)
        for j in range(len(cols)): r.cells[off + j].width = Cm(cell_w)
    no_split(t)
    spacer()
    return t


PLACE = ["一", "十", "百", "千", "萬", "十萬", "百萬", "千萬"]


def vgrid(n, rows, cw=0.95, h=0.85, size=18, header=True, center=True, parent=None):
    """直式格子（n 個數字欄＋左邊運算符號欄）。rows：dict 清單
       d：數字字串（靠右對齊，空白＝空格）；op：運算符號；line：這列下面畫粗線；
       carry：True＝進位小格列；ex：示範藍字；hl：要塗色的位（從右數 0＝一位）；pre：{位: 字} 灰底提示字
       start：k＝乘以二位數的第 2 列「從第 k 位開始寫」：這一列的第 k 位和乘數（op ×）那一列的第 k 位塗同一個橘色。
       老師 2026-10-15：沒有教「第 2 列先寫 0」，教的是「乘哪一位，放下來的就跟它對齊」→ 不要用灰底 0。"""
    rows = [dict(r) for r in rows]
    for r in rows:                       # 舊寫法 pre={0: "0"}（第 2 列先寫 0）一律改成對齊十位
        if r.get("pre") == {0: "0"}:
            del r["pre"]; r["start"] = 1
    starts = {r["start"] for r in rows if "start" in r}
    for r in rows:
        if "start" in r:
            r["hl"] = set(r.get("hl", ())) | {r["start"]}
        elif r.get("op") == "×" and starts:
            r["hl"] = set(r.get("hl", ())) | starts
    t = (parent or doc).add_table(rows=len(rows) + (1 if header else 0), cols=n + 1)
    if center: t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, sz=4, color="BFBFBF")
    r0 = 0
    if header:
        cell_text(t.cell(0, 0), "")
        for j in range(n):
            k = n - 1 - j
            cell_text(t.cell(0, j + 1), PLACE[k], size=9, fill=WAN if k >= 4 else GE)
        row_h(t.rows[0], 0.5)
        r0 = 1
    for i, r in enumerate(rows):
        ri = r0 + i
        d = (r.get("d") or "").rjust(n)
        cell_text(t.cell(ri, 0), r.get("op", ""), size=size, bold=True)
        for j in range(n):
            k = n - 1 - j
            ch = d[j].strip()
            fill = None
            if k in r.get("hl", ()): fill = HOT
            if r.get("carry"): fill = fill or "F7F7F7"
            col = EX if r.get("ex") else None
            if k in r.get("pre", {}):
                ch = r["pre"][k]; fill = "E7E6E6"; col = RGBColor(0x80, 0x80, 0x80)
            cell_text(t.cell(ri, j + 1), ch, size=11 if r.get("carry") else size, color=col, fill=fill)
        row_h(t.rows[ri], 0.55 if r.get("carry") else h)
        if size <= 15: t.rows[ri].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        if r.get("line"):
            for j in range(n + 1):
                cell_border(t.cell(ri, j), bottom=(18, "000000"))
    for row in t.rows:
        row.cells[0].width = Cm(0.9)
        for j in range(n): row.cells[j + 1].width = Cm(cw)
    no_split(t)
    if parent is None: spacer()
    return t


def ldiv(divisor, dividend, q="", work=(), blank=0, cw=0.95, h=0.8, size=18, ex=False, center=True, parent=None):
    """除法直式。work：[(字串, 下面畫線)]，字串長度＝被除數位數（用空白對齊）；blank：多留幾列空白"""
    n = len(dividend)
    rows = 2 + len(work) + blank
    t = (parent or doc).add_table(rows=rows, cols=n + 1)
    if center: t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, sz=0)
    col = EX if ex else None
    qq = q.rjust(n) if q else " " * n
    cell_text(t.cell(0, 0), "")
    for j in range(n):
        cell_text(t.cell(0, j + 1), qq[j].strip(), size=size, color=col, fill=None if q else "F7F7F7")
        cell_border(t.cell(0, j + 1), left=(4, "D9D9D9"), right=(4, "D9D9D9"), top=(4, "D9D9D9"))
    cell_text(t.cell(1, 0), divisor, size=size)
    cell_border(t.cell(1, 0), right=(18, "000000"))
    for j in range(n):
        cell_text(t.cell(1, j + 1), dividend[j], size=size)
        cell_border(t.cell(1, j + 1), top=(18, "000000"))
    for i, (s, line) in enumerate(work):
        s = s.rjust(n) if len(s) < n else s
        for j in range(n):
            cell_text(t.cell(2 + i, j + 1), s[j].strip(), size=size, color=col)
            cell_border(t.cell(2 + i, j + 1), left=(4, "E7E6E6"), right=(4, "E7E6E6"))
            if line: cell_border(t.cell(2 + i, j + 1), bottom=(14, "000000"))
    for i in range(blank):
        for j in range(n):
            cell_border(t.cell(2 + len(work) + i, j + 1), left=(4, "E7E6E6"), right=(4, "E7E6E6"),
                        bottom=(4, "E7E6E6"))
    for r in range(rows):
        row_h(t.rows[r], h)
        if size <= 15 and r >= 2: t.rows[r].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    for row in t.rows:
        row.cells[0].width = Cm(1.3 if len(divisor) > 1 else 0.9)
        for j in range(n): row.cells[j + 1].width = Cm(cw)
    no_split(t)
    if parent is None: spacer()
    return t


def hint(*lines, title="小提醒"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, sz=6, color="A6A6A6")
    c = t.cell(0, 0); c.width = Cm(17.2)
    shade(c._tc.get_or_add_tcPr(), "F2F2F2")
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
    run(p, "★ " + title, bold=True, size=12)
    for ln in lines:
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(1)
        q.paragraph_format.line_spacing = 1.15
        _parts(q, ln if isinstance(ln, (list, tuple)) else [ln], size=12)
    spacer()


def title(part):
    t = doc.add_table(rows=1, cols=2)
    borders(t, sz=0)
    a, b = t.cell(0, 0), t.cell(0, 1)
    a.width, b.width = Cm(11.4), Cm(6.0)
    p = a.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    run(p, HEADER, size=10, color=RGBColor(0x80, 0x80, 0x80))
    q = a.add_paragraph(); q.paragraph_format.space_after = Pt(0)
    run(q, part, bold=True, size=17)
    r = b.paragraphs[0]; r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(r, "座號：____　姓名：________", size=11)
    spacer()


def step(n, *parts, indent=0.4, before=2):
    return para(("第" + str(n) + "步　", "b"), *parts, indent=indent, before=before)


def q(label, *parts, before=6):
    return para((label + "　", "b"), *parts, before=before, keep=True)


def blocks(*groups, indent=0.6):
    """積木：groups＝[("10萬", 3), ...]，每種一列"""
    for name, k in groups:
        para(*sum([[(name, "blk"), " "] for _ in range(k)], []), indent=indent)


def image(path_w, scale=1.0, align="c"):
    path, w = path_w
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(w * scale))
    return p


def grid(cells, cols=2, widths=None, h=None, border=False, size=13):
    """把圖和文字排成表格。cells：每格是 list，元素為 ('img', (path,w)) 或 ('p', parts...) 或 ('sp',)"""
    rows = (len(cells) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, sz=4 if border else 0, color="BFBFBF")
    widths = widths or [17.2 / cols] * cols
    for k, items in enumerate(cells):
        c = t.cell(k // cols, k % cols)
        c.width = Cm(widths[k % cols])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        first = True
        for it in items:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(2)
            if it[0] == "img":
                path, w = it[1]
                sc = it[2] if len(it) > 2 else 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Cm(w * sc))
            else:
                _parts(p, it[1:], size=size)
    for r in t.rows:
        if h: row_h(r, h)
        for j, c in enumerate(r.cells): c.width = Cm(widths[j])
    no_split(t)
    spacer()
    return t


def boxes(n, labels=None, w=1.3, h=1.0):
    """n 個空格（數字卡、寫答案用），labels 寫在上面"""
    t = doc.add_table(rows=2 if labels else 1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    if labels:
        for j, s in enumerate(labels):
            cell_text(t.cell(0, j), s, size=11, bold=True, fill=WAN if "萬" in s else GE)
        row_h(t.rows[0], 0.7)
    row_h(t.rows[-1], h)
    for r in t.rows:
        for c in r.cells: c.width = Cm(w)
    no_split(t)
    spacer()
    return t


def free_app(text, lines=6, size=13):
    """應用題（學生自己列式、自己寫直式）：給 side() 用的 (標題, builder)。
    老師 2026-10-15：不印算式框和直式格、「答」後面不印單位，讓學生自己寫。"""
    def build(c):
        for _ in range(lines):
            e = c.add_paragraph(); e.paragraph_format.space_after = Pt(2)
        a = c.add_paragraph(); a.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(a, "答：＿＿＿＿", size=size)
    return ([text, "\n列式：", "\n直式："], build)


def side(items, widths=None, size=13):
    """並排：items＝[(標題 parts, builder(cell)), ...]"""
    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, sz=0)
    widths = widths or [17.2 / len(items)] * len(items)
    for k, (parts, builder) in enumerate(items):
        c = t.cell(0, k); c.width = Cm(widths[k])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        _parts(p, parts, size=size)
        if builder: builder(c)
    trPr = t.rows[0]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))   # 並排的兩題不要被切到兩頁
    spacer()
    return t


def mults(divisor, n, filled=False, w=1.45):
    """倍數表：除數×1～×n"""
    t = doc.add_table(rows=2, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    d = int(divisor)
    for j in range(n):
        cell_text(t.cell(0, j), f"{divisor}×{j + 1}", size=11, bold=True, fill=GE)
        cell_text(t.cell(1, j), str(d * (j + 1)) if filled else "", size=15, color=EX if filled else None)
    row_h(t.rows[1], 0.85)
    for r in t.rows:
        for c in r.cells: c.width = Cm(w)
    no_split(t)
    spacer()
    return t